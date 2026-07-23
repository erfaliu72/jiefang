import paramiko, time, json

host = '8.156.91.120'
password = 'fijSec-tuccyt-9dohxu'

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect(host, username='root', password=password, timeout=15)

# 在服务器写测试脚本
init_script = """
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
headers = ["序号","经销商代码","经销商名称","网络类型","提车地点","产品代码","产品名称","公告型号","技术路线","能源类型","厂商指导价","开票价格","产品类别","VIN","发动机号","结算价格","资金来源","合格证号","入库日期","提车仓库","驾驶室","驱动形式","发动机厂家","发动机功率","变速箱","后桥","轴距","轮胎","速比"]
for i, h in enumerate(headers):
    ws.cell(row=3, column=i+1, value=h)
ws.cell(row=4, column=14, value="LFNA4LDC6TAE31340")
ws.cell(row=4, column=15, value="ENG123")
ws.cell(row=4, column=7, value="测试车型")
ws.cell(row=4, column=6, value="PROD001")
wb.save("/tmp/vintest.xlsx")
print("Excel created OK")
"""
_, stdout, _ = ssh.exec_command('cat > /tmp/make_test.py && chmod +x /tmp/make_test.py')
print('Script placeholder created')

sftp = ssh.open_sftp()
with sftp.open('/tmp/make_test.py', 'w') as f:
    f.write(init_script)
sftp.close()

stdin, stdout, stderr = ssh.exec_command('/opt/jinjuyuan-venv/bin/python3 /tmp/make_test.py')
print('Create Excel:', stdout.read().decode().strip())

# Login
stdin, stdout, stderr = ssh.exec_command(
    'curl -s -c /tmp/fleet_cookie -X POST '
    '-H "Content-Type: application/json" '
    '-d \'{"username":"fleet","password":"123456"}\' '
    'http://127.0.0.1:8081/api/auth/login'
)
login_resp = stdout.read().decode()
print('Login:', login_resp[:200])

# Upload test
stdin, stdout, stderr = ssh.exec_command(
    'curl -s -b /tmp/fleet_cookie '
    '-X POST -F "file=@/tmp/vintest.xlsx" '
    'http://127.0.0.1:8081/api/vehicles/import'
)
import_resp = stdout.read().decode()
print('Import response:', import_resp[:500])
try:
    d = json.loads(import_resp)
    print('VIN details:', d.get('details', []))
except:
    pass

ssh.close()
